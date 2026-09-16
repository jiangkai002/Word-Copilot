# -*- coding: utf-8 -*-
"""
生成测试文档（需求文档 §67）：

    test-documents/
      01-simple-paragraph.docx    单段落：基础 Diff / 修订测试
      02-heading-document.docx    多级标题：大纲 / 章节上下文测试
      03-table.docx               表格文档：复杂内容边界测试
      04-existing-revisions.docx  含人工修订（Revision A）：§69 修订隔离测试
      05-formatting.docx          富格式（加粗/斜体/颜色/字号）：格式保持测试
      06-long-document.docx       超长文档（>5 万字符）：全文模式截断回退测试
      07-chinese-document.docx    中文标点/引号/混排：重点中文测试
      08-english-document.docx    英文文档

用法（在仓库根目录）：

    py -m pip install python-docx
    py tools/generate_test_documents.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "test-documents"

MANUAL_AUTHOR = "手工修订用户"
REVISION_DATE = "2026-09-15T10:00:00Z"


# ---------------------------------------------------------------------------
# 修订（Track Changes）XML 工具：python-docx 本身不暴露修订 API，
# 直接在 OOXML 层构造 w:ins / w:del。
# ---------------------------------------------------------------------------

def wrap_run_in_insertion(run, rev_id: int, author: str = MANUAL_AUTHOR) -> None:
    """把已有 run 包进 w:ins —— Word 中显示为「插入的修订」。"""
    r = run._r
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), str(rev_id))
    ins.set(qn("w:author"), author)
    ins.set(qn("w:date"), REVISION_DATE)
    r.addprevious(ins)
    ins.append(r)


def append_deleted_text(paragraph, text: str, rev_id: int, author: str = MANUAL_AUTHOR) -> None:
    """追加 w:del + w:delText —— Word 中显示为「删除的修订」。"""
    delete = OxmlElement("w:del")
    delete.set(qn("w:id"), str(rev_id))
    delete.set(qn("w:author"), author)
    delete.set(qn("w:date"), REVISION_DATE)
    run = OxmlElement("w:r")
    del_text = OxmlElement("w:delText")
    del_text.set(qn("xml:space"), "preserve")
    del_text.text = text
    run.append(del_text)
    delete.append(run)
    paragraph._p.append(delete)


# ---------------------------------------------------------------------------
# 各测试文档
# ---------------------------------------------------------------------------

def build_01_simple_paragraph() -> Document:
    doc = Document()
    doc.add_heading("简单段落测试文档", level=0)
    doc.add_paragraph("系统采用传统架构，可以处理大量数据。")
    doc.add_paragraph("本段用于测试段落模式下的上下文读取与修改。")
    doc.add_paragraph("光标放在这一段上时，AI 应只读取本段及其前后段。")
    return doc


def build_02_heading_document() -> Document:
    doc = Document()
    doc.add_heading("项目技术方案", level=0)

    doc.add_heading("一、背景", level=1)
    doc.add_paragraph(
        "随着业务规模持续扩大，现有单体系统的维护成本逐年上升，"
        "新功能交付周期变长，团队协作冲突频繁。"
    )

    doc.add_heading("二、系统架构", level=1)
    doc.add_heading("2.1 总体架构", level=2)
    doc.add_paragraph(
        "系统整体划分为接入层、业务服务层与数据层三部分，"
        "各层之间通过统一网关通信。"
    )
    doc.add_heading("2.2 数据流程", level=2)
    doc.add_paragraph(
        "用户请求经过接入层鉴权后进入业务服务层，处理结果先写入消息队列，"
        "再由异步任务落库，保证峰值流量下的写入稳定性。"
    )
    doc.add_heading("三、总结", level=1)
    doc.add_paragraph("本方案通过拆分边界与异步化改造，降低耦合并提升整体吞吐能力。")
    return doc


def build_03_table() -> Document:
    doc = Document()
    doc.add_heading("表格测试文档", level=0)
    doc.add_paragraph("下表列出系统当前的数据规模，供容量评估参考。")

    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "数据类型"
    header[1].text = "数量"
    header[2].text = "说明"
    rows = [
        ("订单记录", "约 2000 万", "按月分区存储"),
        ("用户档案", "约 500 万", "含历史手机号"),
        ("日志数据", "约 12 亿", "仅保留 90 天"),
    ]
    for row, values in zip(table.rows[1:], rows, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value

    doc.add_paragraph("表格上方的段落可用于选区测试。")
    doc.add_paragraph("表格下方的段落同样可用于普通文本修改。")
    return doc


def build_04_existing_revisions() -> Document:
    """§69 关键文档：先有人工 Revision A，再让 AI 生成 Revision B。

    验收标准：在插件中接受 AI 修改后，人工 Revision A 必须仍然存在。
    """
    doc = Document()
    doc.add_heading("已有修订测试文档", level=0)

    doc.add_paragraph("本测试文档包含人工创建的修订（Revision A）。")

    # 段落 2：插入修订 + 删除修订
    para = doc.add_paragraph()
    para.add_run("本段包含")
    inserted = para.add_run("人工插入的修订文字，")
    para.add_run("以及")
    wrap_run_in_insertion(inserted, rev_id=1001)
    append_deleted_text(para, "已被人工删除的旧表述，", rev_id=1002)
    para.add_run("结尾保持不变。")

    # 段落 3：干净段落 —— AI 的修改目标（Revision B 落在这里）
    doc.add_paragraph("系统采用传统架构，可以处理大量数据。")

    # 段落 4：第二个干净目标，用于测试同一文档内的多个 AI 事务
    doc.add_paragraph("本段文字与上一段相互独立，用于验证多事务隔离。")
    return doc


def build_05_formatting() -> Document:
    doc = Document()
    doc.add_heading("富格式测试文档", level=0)

    para = doc.add_paragraph()
    run = para.add_run("加粗文字")
    run.bold = True
    para.add_run("与")
    run = para.add_run("斜体文字")
    run.italic = True
    para.add_run("相邻，")
    run = para.add_run("下划线文字")
    run.underline = True
    para.add_run("居中，")
    run = para.add_run("红色文字")
    run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    para.add_run("靠后，")
    run = para.add_run("大字号文字")
    run.font.size = Pt(18)
    para.add_run("收尾。")

    doc.add_heading("带样式的标题段落", level=1)
    doc.add_paragraph("AI 修改本段时，不应破坏上下文的加粗、标题等既有格式。")
    doc.add_paragraph(
        "系统采用传统架构，可以处理大量数据。字符级 Diff 只应触及变化的文字。"
    )
    return doc


_LONG_SECTIONS = [
    ("性能压测报告", "压测环境与结论"),
    ("容量规划说明", "数据增长模型"),
    ("灰度发布方案", "批次与回滚策略"),
    ("安全审计记录", "风险项与整改"),
    ("监控告警治理", "指标与阈值"),
    ("数据迁移计划", "双写与校验"),
    ("接口契约变更", "兼容性说明"),
    ("故障复盘汇总", "根因与改进"),
    ("依赖组件清单", "版本与许可"),
    ("成本优化分析", "资源利用率"),
    ("团队协作规范", "评审与合并"),
    ("发布检查清单", "上线前核对"),
]

_LONG_PARAGRAPHS = [
    "本次评估覆盖核心链路的主要指标，采样周期为连续七个自然日，"
    "统计口径与前次保持一致，异常样本已在报告中单独标注。",
    "按当前增长曲线外推，主存储容量将在两个季度内达到上限，"
    "建议提前启动分片扩容评估，并把冷数据归档纳入例行任务。",
    "灰度批次之间保留至少四十八小时观察窗口，任一批次出现关键告警"
    "即暂停推进并按预案回滚，回滚演练已在预发环境完成。",
    "审计共发现若干待整改项，其中高危项已当场修复，"
    "其余事项均排入下个迭代，责任人明确、期限明确。",
]


def build_06_long_document() -> Document:
    doc = Document()
    doc.add_heading("超长文档测试", level=0)
    doc.add_paragraph(
        "本文档超过 5 万字符，用于验证 @document 模式下的"
        "超长回退（标题结构 + 相关段落）。"
    )

    paragraph_index = 0
    for section_index, (title, subtitle) in enumerate(_LONG_SECTIONS):
        doc.add_heading(f"{section_index + 1}. {title}", level=1)
        doc.add_heading(f"{section_index + 1}.1 {subtitle}", level=2)
        for _ in range(60):
            text = _LONG_PARAGRAPHS[paragraph_index % len(_LONG_PARAGRAPHS)]
            filler = _LONG_PARAGRAPHS[(paragraph_index + 1) % len(_LONG_PARAGRAPHS)]
            paragraph_index += 1
            # 拼接两段填充句 + 编号，把总字符量推到 5 万以上
            doc.add_paragraph(
                text + filler + f"（编号 {paragraph_index:04d}）"
                "该记录用于冲量，内容本身无实际含义，仅用于验证超长文档的回退策略。"
            )
    return doc


def build_07_chinese_document() -> Document:
    doc = Document()
    doc.add_heading("中文场景测试文档", level=0)

    doc.add_paragraph(
        "系统采用传统架构，可以处理大量数据；"
        "「引号」与『书名号』、省略号……以及破折号——都会出现在正式文本中。"
    )
    doc.add_paragraph(
        "验收标准（第一版）：1. 错别字纠正；2. 语病修复；"
        "3. 标点规范化；4. 不改变原意与段落结构。"
    )
    doc.add_paragraph(
        "全角符号（（））、顿号「、」与英文逗号,混排 123、456，"
        "以及 3.14 与 98% 这类数字，都是 Diff 需要正确处理的字符。"
    )
    doc.add_paragraph("这段话里有一个错白字和一个明显地的语病，用于测试纠错命令。")
    doc.add_paragraph(
        "标点后多余空格 、句末双标点。。连续的逗号，，，"
        "都是常见的中文文本文本问题。"
    )
    return doc


def build_08_english_document() -> Document:
    doc = Document()
    doc.add_heading("English Test Document", level=0)

    doc.add_paragraph(
        "The system uses a traditional architecture and handles "
        "a large volume of data every day."
    )
    doc.add_paragraph(
        "This paragraph contains several deliberate mistkes, "
        "including speling errors and awkward phrasing that need correction."
    )
    doc.add_paragraph(
        "Key milestones for the migration: database dual-write, "
        "backfill verification, and traffic cutover."
    )
    doc.add_paragraph(
        "Another independent paragraph used to verify that multiple "
        "edit transactions remain isolated from each other."
    )
    return doc


BUILDERS = [
    ("01-simple-paragraph.docx", build_01_simple_paragraph),
    ("02-heading-document.docx", build_02_heading_document),
    ("03-table.docx", build_03_table),
    ("04-existing-revisions.docx", build_04_existing_revisions),
    ("05-formatting.docx", build_05_formatting),
    ("06-long-document.docx", build_06_long_document),
    ("07-chinese-document.docx", build_07_chinese_document),
    ("08-english-document.docx", build_08_english_document),
]


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for filename, builder in BUILDERS:
        doc = builder()
        target = OUTPUT_DIR / filename
        doc.save(str(target))
        print(f"已生成 {target}")
    print(f"\n全部 {len(BUILDERS)} 个测试文档已生成于 {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
